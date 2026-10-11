"""Capture and extract original research/regulator documents without certifying findings."""
from __future__ import annotations

import argparse
import hashlib
import ipaddress
import io
import json
import re
import socket
import sys
import urllib.request
from pathlib import Path
from urllib.parse import parse_qsl, urlsplit
from xml.etree import ElementTree

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'scripts/validate_apis'))
from common import get_secret, now_iso, redact_sensitive, with_query
from scripts.market_research.research_sources import public_url


def public_address(url):
    parts = urlsplit(url)
    if parts.scheme != 'https' or not parts.hostname or parts.username or parts.password:
        raise ValueError('Use a public HTTPS source URL without embedded credentials')
    addresses = socket.getaddrinfo(parts.hostname, parts.port or 443, type=socket.SOCK_STREAM)
    if any(not ipaddress.ip_address(row[4][0]).is_global for row in addresses):
        raise ValueError('Private/local network document URLs are not supported')
    return parts


class PublicRedirects(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        public_address(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def download(url, maximum=10000000):
    public_address(url)
    request = urllib.request.Request(url, headers={'User-Agent': 'business-strategist-research/1.0'})
    with urllib.request.build_opener(PublicRedirects()).open(request, timeout=30) as response:
        raw = response.read(maximum + 1)
        if len(raw) > maximum:
            raise ValueError('Document exceeds byte budget; narrow the requested material')
        return raw, response.headers.get('Content-Type', '').split(';', 1)[0]


def extract(raw, kind, limit=100):
    if kind == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError('Encrypted PDFs need an accessible original version')
        texts = [(str(index + 1), page.extract_text() or '') for index, page in enumerate(reader.pages[:limit])]
        complete = len(reader.pages) <= limit
    elif kind == 'xml':
        text = raw.decode('utf-8')
        if '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper():
            raise ValueError('Document type/entity declarations are not supported')
        root = ElementTree.fromstring(text)
        texts = [('document', ' '.join(root.itertext()))]
        complete = True
    elif kind == 'html':
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(raw, 'html.parser')
        for node in soup(['script', 'style', 'nav']):
            node.decompose()
        texts = [('document', soup.get_text(' ', strip=True))]
        complete = True
    elif kind == 'xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        texts = []
        complete = True
        for sheet in workbook:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= limit:
                    complete = False
                    break
                rows.append(json.dumps(list(row), default=str, ensure_ascii=False))
            texts.append((sheet.title, '\n'.join(rows)))
        workbook.close()
    else:
        texts, complete = [('document', raw.decode('utf-8-sig'))], True
    if sum(len(text) for _, text in texts) > 2000000:
        raise ValueError('Extracted text exceeds review budget')
    if not any(text.strip() for _, text in texts):
        raise ValueError('No readable text extracted; scanned PDF may require OCR')
    return texts, complete


def main():
    p = argparse.ArgumentParser(description=__doc__)
    route = p.add_mutually_exclusive_group(required=True)
    route.add_argument('--url')
    route.add_argument('--openalex-work', help='Explicit W... ID; downloads TEI XML with the existing OpenAlex key')
    route.add_argument('--input', type=Path, help='Previously obtained local original document')
    p.add_argument('--source-url', help='Original public URL, required for local input')
    p.add_argument('--publisher', required=True)
    p.add_argument('--source-role', choices=['original_research', 'regulator', 'company_register'], default='original_research')
    p.add_argument('--format', choices=['pdf', 'xml', 'html', 'xlsx', 'csv', 'text'])
    p.add_argument('--limit', type=int, default=100, help='Maximum PDF pages or rows per spreadsheet sheet')
    p.add_argument('--out', type=Path, required=True)
    args = p.parse_args()
    if args.out.exists() or not 1 <= args.limit <= 1000:
        p.error('Use a fresh output directory and limit 1–1000')
    try:
        if args.openalex_work:
            if not re.fullmatch(r'W\d+', args.openalex_work):
                raise ValueError('OpenAlex work must be an explicit W... identifier')
            key = get_secret('OPENALEX_API_KEY')[1]
            if not key:
                raise ValueError('Missing OPENALEX_API_KEY')
            source = 'https://content.openalex.org/works/' + args.openalex_work + '.grobid-xml'
            raw, mime = download(with_query(source, {'api_key': key}))
            kind = 'xml'
        elif args.input:
            if not args.source_url:
                raise ValueError('--source-url is required for local originals')
            source, mime = args.source_url, ''
            if urlsplit(source).scheme != 'https':
                raise ValueError('Original document requires an HTTPS source URL')
            if args.input.stat().st_size > 10000000:
                raise ValueError('Original exceeds byte budget')
            raw = args.input.read_bytes()
            kind = args.format or args.input.suffix.lower().lstrip('.')
        else:
            source = args.url
            if any(key.lower() in {'api_key', 'key', 'token', 'access_token'} for key, _ in parse_qsl(urlsplit(source).query)):
                raise ValueError('Use --openalex-work for authenticated content; public URLs must not embed credentials')
            raw, mime = download(source)
            suffix = Path(urlsplit(source).path).suffix.lower().lstrip('.')
            kind = args.format or {'application/pdf': 'pdf', 'text/html': 'html', 'text/xml': 'xml', 'application/xml': 'xml',
                                   'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': 'xlsx'}.get(mime) or suffix
        if kind not in {'pdf', 'xml', 'html', 'xlsx', 'csv', 'text'}:
            raise ValueError('Unsupported document format; supply --format for a verified original')
        texts, complete = extract(raw, kind, args.limit)
        args.out.mkdir(parents=True)
        original = args.out / ('original.' + kind)
        original.write_bytes(raw)
        observations = [{'text': text, 'original_locator': locator, 'currency': None, 'scope': 'Unreviewed original document'} for locator, text in texts]
        capture = {'provider': 'original_document', 'publisher': args.publisher, 'source_url': public_url(source),
                   'source_role': args.source_role, 'retrieved_at': now_iso(), 'original_path': original.name,
                   'original_sha256': hashlib.sha256(raw).hexdigest(), 'scope': {'population': 'Unreviewed original document'},
                   'observations': observations, 'extraction_complete': complete, 'review_status': 'pending'}
        (args.out / 'capture.json').write_text(json.dumps(capture, indent=2, ensure_ascii=False) + '\n')
        capture_digest = hashlib.sha256((args.out / 'capture.json').read_bytes()).hexdigest()
        facts = []
        for index, observation in enumerate(observations):
            ident = capture['original_sha256'] + ':' + observation['original_locator']
            facts.append({'schema_version': '1.0', 'evidence_id': 'fact-' + hashlib.sha256(ident.encode()).hexdigest()[:20],
                          'fact_kind': 'company_observation' if args.source_role == 'company_register' else 'study_finding',
                          'source': 'original_document', 'source_record_id': ident, 'source_url': capture['source_url'],
                          'publisher': args.publisher, 'source_role': args.source_role, 'retrieved_at': capture['retrieved_at'],
                          'reference_period': 'unknown', 'geography': 'unknown', 'population': capture['scope']['population'],
                          'original_language': 'unknown', 'capture_path': 'capture.json', 'content_sha256': capture_digest,
                          'locator': '/observations/' + str(index), 'observation': observation, 'text': observation['text'],
                          'review_status': 'pending', 'uncertainty': ['Original passage candidate; no substantive finding is certified until methods, context and applicability are reviewed.']})
        (args.out / 'research-facts.jsonl').write_text(''.join(json.dumps(row, ensure_ascii=False) + '\n' for row in facts))
        (args.out / 'extracted.md').write_text('\n\n'.join('## ' + locator + '\n\n' + text for locator, text in texts))
        print(json.dumps({'status': 'ok', 'output': str(args.out), 'extraction_complete': complete, 'review_status': 'pending'}))
    except Exception as exc:
        p.exit(2, str(redact_sensitive(str(exc))) + '\n')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
